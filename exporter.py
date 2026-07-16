"""Генерира kontakti_znamena.xlsx от JSON checkpoint-ите в data/.

Листове: "Всички" (дедуплицирани), по един за всяка категория, "Статистика".
Форматиране: замразен header, автофилтри, автоширина, header с фон.
"""
import json
import logging
from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from base import normalize_phone, validate_email

log = logging.getLogger("exporter")

COLUMNS = ["kategoria", "podkategoria", "ime", "grad", "adres", "telefon",
           "email", "uebsait", "dopalnitelno", "iztochnik", "data_izvlichane"]

HEADER_BG = ["Категория", "Подкатегория", "Име", "Град", "Адрес", "Телефон",
             "Имейл", "Уебсайт", "Допълнително", "Източник", "Дата"]


def load_all_records() -> list[dict]:
    records = []
    for path in sorted(config.DATA_DIR.glob("*.json")):
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            log.warning("Невалиден JSON: %s — пропускам", path.name)
            continue
        recs = payload.get("records", payload if isinstance(payload, list) else [])
        log.info("%s: %d записа", path.name, len(recs))
        records.extend(recs)
    return records


def validate_records(records: list[dict]) -> list[dict]:
    """Валидира имейли/телефони; невалидните стойности се изчистват."""
    out = []
    for r in records:
        r = {c: str(r.get(c, "") or "").strip() for c in COLUMNS}
        if not r["ime"]:
            continue
        r["email"] = validate_email(r["email"]) or ""
        r["telefon"] = normalize_phone(r["telefon"]) or ""
        out.append(r)
    return out


def dedupe_records(records: list[dict]) -> list[dict]:
    """Дедупликация по нормализиран телефон И по имейл.

    Един запис може да идва от два каталога — при съвпадение на имейл или
    телефон записите се сливат, като празните полета се допълват.
    """
    result: list[dict] = []
    by_email: dict[str, dict] = {}
    by_phone: dict[str, dict] = {}
    by_name: dict[tuple, dict] = {}

    for r in records:
        existing = None
        if r["email"] and r["email"] in by_email:
            existing = by_email[r["email"]]
        elif r["telefon"] and r["telefon"] in by_phone:
            existing = by_phone[r["telefon"]]
        elif not r["email"] and not r["telefon"]:
            key = (r["ime"].lower(), r["grad"].lower(), r["kategoria"])
            existing = by_name.get(key)
            if existing is None:
                by_name[key] = r

        if existing is not None:
            for c in COLUMNS:  # допълваме празните полета от новия запис
                if not existing[c] and r[c]:
                    existing[c] = r[c]
        else:
            result.append(r)

        # Индексите сочат към живия (запазен или слят) запис.
        final = existing if existing is not None else r
        if final["email"]:
            by_email.setdefault(final["email"], final)
        if final["telefon"]:
            by_phone.setdefault(final["telefon"], final)

    return result


def _style_sheet(ws, ncols: int, nrows: int) -> None:
    header_fill = PatternFill("solid", start_color="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(nrows, 1) + 1}"
    # Автоматична ширина по най-дългата стойност (с таван).
    for col_idx in range(1, ncols + 1):
        letter = get_column_letter(col_idx)
        longest = max(
            (len(str(c.value)) for c in ws[letter] if c.value is not None),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 60)


def export_excel(records: list[dict] | None = None) -> "pd.DataFrame":
    if records is None:
        records = load_all_records()
    total_raw = len(records)
    records = validate_records(records)
    records = dedupe_records(records)
    log.info("Записи: %d сурови -> %d след валидация и дедупликация",
             total_raw, len(records))

    df = pd.DataFrame(records, columns=COLUMNS)
    config.OUTPUT_DIR.mkdir(exist_ok=True)

    with pd.ExcelWriter(config.EXCEL_PATH, engine="openpyxl") as writer:
        def write_sheet(name: str, frame: pd.DataFrame):
            out = frame.copy()
            out.columns = HEADER_BG
            out.to_excel(writer, sheet_name=name[:31], index=False)
            _style_sheet(writer.sheets[name[:31]], len(COLUMNS), len(out))

        write_sheet("Всички", df)

        for kategoria, sheet in config.CATEGORY_SHEETS.items():
            sub = df[df["kategoria"] == kategoria]
            if sub.empty:
                continue
            if sheet[:31] in writer.sheets:  # две категории в общ лист
                startrow = writer.sheets[sheet[:31]].max_row
                out = sub.copy()
                out.columns = HEADER_BG
                out.to_excel(writer, sheet_name=sheet[:31], index=False,
                             header=False, startrow=startrow)
                _style_sheet(writer.sheets[sheet[:31]], len(COLUMNS),
                             writer.sheets[sheet[:31]].max_row - 1)
            else:
                write_sheet(sheet, sub)

        # Лист "Статистика".
        stats_rows = []
        for kategoria in df["kategoria"].unique():
            sub = df[df["kategoria"] == kategoria]
            stats_rows.append({
                "Категория": kategoria,
                "Брой записи": len(sub),
                "% с имейл": round(100 * (sub["email"] != "").mean(), 1),
                "% с телефон": round(100 * (sub["telefon"] != "").mean(), 1),
            })
        stats_rows.append({
            "Категория": "ОБЩО",
            "Брой записи": len(df),
            "% с имейл": round(100 * (df["email"] != "").mean(), 1) if len(df) else 0,
            "% с телефон": round(100 * (df["telefon"] != "").mean(), 1) if len(df) else 0,
        })
        stats = pd.DataFrame(stats_rows)
        stats["Дата на генериране"] = date.today().isoformat()
        stats.to_excel(writer, sheet_name="Статистика", index=False)
        _style_sheet(writer.sheets["Статистика"], len(stats.columns), len(stats))

    log.info("Excel записан: %s (%d записа)", config.EXCEL_PATH, len(df))
    return df
