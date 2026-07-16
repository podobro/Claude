"""Тестове на пайплайна: нормализация, валидация, дедупликация, Excel.

Стартиране:  python -m pytest tests/ -v   (или  python tests/test_pipeline.py)
"""
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import config
from exporter import dedupe_records, export_excel, validate_records
from base import extract_emails, extract_phones, normalize_phone, validate_email


def rec(**kw):
    base = {c: "" for c in ("kategoria", "podkategoria", "ime", "grad", "adres",
                            "telefon", "email", "uebsait", "dopalnitelno",
                            "iztochnik")}
    base["kategoria"] = "Общини"
    base["data_izvlichane"] = date.today().isoformat()
    base.update(kw)
    return base


def test_normalize_phone():
    assert normalize_phone("032 656 701") == "+35932656701"
    assert normalize_phone("+359 32 656-701") == "+35932656701"
    assert normalize_phone("00359 2 9434 467") == "+35929434467"
    assert normalize_phone("0887476017") == "+359887476017"
    assert normalize_phone("+49 30 201 090") == "+4930201090"  # чужд номер остава
    assert normalize_phone("123") is None
    assert normalize_phone("") is None


def test_validate_email():
    assert validate_email("Info@Plovdiv.BG") == "info@plovdiv.bg"
    assert validate_email("не-е-имейл") is None
    assert validate_email("a@b") is None


def test_extract_from_text():
    text = ("Община Пловдив, пл. Стефан Стамболов 1, тел: 032/656 701, "
            "e-mail: info@plovdiv.bg, ivan.petrov@gmail.com")
    assert "info@plovdiv.bg" in extract_emails(text)
    # Личен gmail без служебен префикс се пропуска.
    assert "ivan.petrov@gmail.com" not in extract_emails(text)
    assert "+35932656701" in extract_phones(text)


def test_dedupe_merges_by_email_and_phone():
    records = validate_records([
        rec(ime="Община Пловдив", email="info@plovdiv.bg", telefon="032656701"),
        # Дубликат по имейл от друг каталог — с допълнителен адрес.
        rec(ime="ОБЩИНА ПЛОВДИВ", email="info@plovdiv.bg",
            adres="пл. Стефан Стамболов 1"),
        # Дубликат по телефон (нормализиран различен запис).
        rec(ime="Municipality Plovdiv", telefon="+359 32 656 701",
            uebsait="https://www.plovdiv.bg"),
        # Различен запис.
        rec(ime="Община Варна", email="info@varna.bg"),
        # Без имейл/телефон — дедуп по (име, град, категория).
        rec(ime="Община Бургас", grad="Бургас"),
        rec(ime="Община Бургас", grad="Бургас", adres="ул. Александровска 26"),
    ])
    result = dedupe_records(records)
    assert len(result) == 3, [r["ime"] for r in result]
    plovdiv = next(r for r in result if "Пловдив" in r["ime"])
    # Слетият запис е допълнен с полетата от дубликатите.
    assert plovdiv["adres"] == "пл. Стефан Стамболов 1"
    assert plovdiv["uebsait"] == "https://www.plovdiv.bg"
    assert plovdiv["telefon"] == "+35932656701"


def test_export_excel(tmp_path=None):
    records = [
        rec(kategoria="Общини", ime="Община Пловдив", grad="Пловдив",
            email="info@plovdiv.bg", telefon="032656701"),
        rec(kategoria="Общини", ime="Община Пловдив", grad="Пловдив",
            email="info@plovdiv.bg"),  # изкуствен дубликат
        rec(kategoria="Посолства", ime="Посолство на РБ в Берлин",
            telefon="+49 30 201 090", email="embassy.berlin@mfa.bg"),
        rec(kategoria="Хотели", ime="Хотел Тримонциум", grad="Пловдив",
            dopalnitelno="4 звезди"),
        rec(kategoria="Бизнес клиенти", podkategoria="Рекламна агенция",
            ime="Агенция Х", telefon="0888123456"),
    ]
    df = export_excel(records)
    assert len(df) == 4  # дубликатът е слят

    from openpyxl import load_workbook
    wb = load_workbook(config.EXCEL_PATH)
    for sheet in ("Всички", "Общини", "Посолства", "Хотели", "Бизнес", "Статистика"):
        assert sheet in wb.sheetnames, f"липсва лист {sheet}"
    ws = wb["Всички"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    stats = wb["Статистика"]
    values = [row[0].value for row in stats.iter_rows(min_row=2)]
    assert "ОБЩО" in values
    print(f"OK: Excel генериран с листове {wb.sheetnames}")


if __name__ == "__main__":
    test_normalize_phone()
    test_validate_email()
    test_extract_from_text()
    test_dedupe_merges_by_email_and_phone()
    test_export_excel()
    print("Всички тестове минаха успешно.")
